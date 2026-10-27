"""Round trips the descriptive report text between the language file and a Word document.

    python Convert-ReportTextWord.py export ReportText.docx
    python Convert-ReportTextWord.py import ReportText.docx

Only prose strings (60 characters or more) are exported. Short labels such as column names are
left in the language file. In the document each translation group is a Heading 1, each key a
Heading 2, and the text is the paragraph beneath it. Edit the paragraphs only: the headings are
how the import finds the string to replace. Keep {0}, {1} style placeholders as they are.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
LANG = ROOT / "AsBuiltReport.Ribbon.SWeEdge" / "Language" / "en-US" / "RibbonSWeEdge.psd1"
MIN_LENGTH = 60


def load_strings():
    # Importing through PowerShell gives the values after ConvertFrom-StringData has unescaped them.
    script = (
        "$d = Import-LocalizedData -BaseDirectory '{0}' -FileName RibbonSWeEdge -UICulture en-US;"
        "$o = [ordered]@{{}}; foreach ($g in $d.Keys) {{ $h = [ordered]@{{}};"
        "foreach ($k in $d[$g].Keys) {{ $h[$k] = [string]$d[$g][$k] }}; $o[$g] = $h }};"
        "$o | ConvertTo-Json -Depth 4"
    ).format(LANG.parent.parent)
    out = subprocess.run(["pwsh", "-NoProfile", "-Command", script], capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


def export(path):
    doc = Document()
    doc.add_heading("Ribbon SWeEdge report text", 0)
    doc.add_paragraph(
        "Edit the paragraphs only. Do not change the headings, and keep placeholders such as {0} and {1}."
    )
    for group, items in load_strings().items():
        prose = {k: v for k, v in items.items() if len(v) >= MIN_LENGTH}
        if not prose:
            continue
        doc.add_heading(group, 1)
        for key, value in prose.items():
            doc.add_heading(key, 2)
            doc.add_paragraph(value)
    doc.save(path)
    print(f"Wrote {path}")


def to_data_line(value):
    # ConvertFrom-StringData treats a backslash as an escape, and each value must be one line.
    return " ".join(value.split()).replace("\\", "\\\\")


def import_text(path):
    doc = Document(path)
    edits, group, key = {}, None, None
    for p in doc.paragraphs:
        style = p.style.name
        if style == "Heading 1":
            group, key = p.text.strip(), None
        elif style == "Heading 2":
            key = p.text.strip()
            edits[(group, key)] = ""
        elif key and p.text.strip():
            edits[(group, key)] = (edits[(group, key)] + " " + p.text).strip()

    current = load_strings()
    lines = LANG.read_text(encoding="utf-8-sig").splitlines(keepends=True)
    changed = 0
    group = None
    for i, line in enumerate(lines):
        m = re.match(r"\s*(\w+) = ConvertFrom-StringData @'", line)
        if m:
            group = m.group(1)
            continue
        m = re.match(r"(\s*)(\w+)(\s*=\s*)(.*?)(\r?\n)?$", line)
        if not (group and m and (group, m.group(2)) in edits):
            continue
        new = edits[(group, m.group(2))]
        if not new or new == current[group][m.group(2)]:
            continue
        lines[i] = f"{m.group(1)}{m.group(2)}{m.group(3)}{to_data_line(new)}{m.group(5) or ''}"
        changed += 1
    LANG.write_text("".join(lines), encoding="utf-8-sig")
    print(f"Updated {changed} string(s) in {LANG.name}")


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("export", "import"):
        sys.exit(__doc__)
    (export if sys.argv[1] == "export" else import_text)(sys.argv[2])
